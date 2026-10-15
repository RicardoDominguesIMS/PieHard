"""Read supported recipe document formats as plain text."""

from html.parser import HTMLParser
from pathlib import Path


SUPPORTED_EXTENSIONS = {".txt", ".md", ".html", ".htm", ".pdf", ".docx"}


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self._ignored_depth += 1
        elif not self._ignored_depth and tag in {"br", "p", "div", "li", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._ignored_depth:
            self._ignored_depth -= 1
        elif not self._ignored_depth and tag in {"p", "div", "li", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth:
            self.parts.append(data)


def read_recipe_units(path: Path, source_path: str) -> list[tuple[str, str]]:
    """Extract source text, using one independently traceable unit per PDF page."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported recipe document type: {suffix or '(no extension)'}")

    if suffix == ".pdf":
        import pymupdf

        units = []
        with pymupdf.open(path) as document:
            for page_number, page in enumerate(document, start=1):
                text = page.get_text("text").strip()
                units.append((f"{source_path}#page={page_number}", text))
        if not units:
            raise ValueError("PDF contains no pages.")
        return units

    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8-sig")
    elif suffix in {".html", ".htm"}:
        parser = _HTMLTextExtractor()
        parser.feed(path.read_text(encoding="utf-8-sig"))
        text = "".join(parser.parts)
    else:
        from docx import Document

        document = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            text += "\n" + "\n".join(
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            )

    text = text.strip()
    if not text:
        raise ValueError("No readable text was found in the document.")
    return [(source_path, text)]


def read_recipe_document(path: Path) -> str:
    """Extract all readable text from a supported recipe document."""
    source_path = path.name
    units = read_recipe_units(path, source_path)
    return "\n\n".join(text for _, text in units if text)
